#!/usr/bin/env python
"""
construir.py — Genera inventario-tareas.xlsx con las fórmulas descritas en inventario-tareas.md.

    python contenido/claude-en-tu-empresa/plantillas/construir.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
NAVY, CYAN, TEXTO = "0A1A2F", "00A0DB", "F5F8FB"

wb = Workbook()

# --- Supuestos ----------------------------------------------------------------
sup = wb.active
sup.title = "Supuestos"
sup.append(["Supuesto", "Valor", "Origen"])
filas = [
    ("Horas por semana", 42.5, "landing, data.js"),
    ("Factor de carga social", 1.35, "landing"),
    ("Semanas al año", 46, "landing"),
    ("Factor de realización", 0.6, "landing"),
    ("Jornada 2027 (h/semana)", 46, "jornada-40"),
    ("Jornada 2030 (h/semana)", 40, "jornada-40"),
    ("Umbral horas/año para candidata", 20, "criterio del curso"),
]
for f in filas:
    sup.append(list(f))
for c in ("A", "B", "C"):
    sup.column_dimensions[c].width = 32
# Nombres: B2 horas/semana, B3 carga, B4 semanas, B5 realización, B8 umbral

# --- Tareas ------------------------------------------------------------------
t = wb.create_sheet("Tareas")
cab = ["Área", "Tarea", "Quién la hace", "Frecuencia", "Veces por semana", "Minutos cada vez",
       "Tipo (A/B)", "Sueldo mensual bruto", "Horas al año", "Costo al año", "Candidata a IA",
       "Primer paso", "Quién revisa"]
t.append(cab)
for i, w in enumerate([16, 34, 18, 13, 16, 16, 12, 20, 14, 16, 15, 22, 18], start=1):
    t.column_dimensions[get_column_letter(i)].width = w

N = 200
for r in range(2, N + 2):
    # Veces por semana derivada de la frecuencia, editable.
    t[f"E{r}"] = f'=IF(D{r}="diaria",5,IF(D{r}="semanal",1,IF(D{r}="quincenal",0.5,IF(D{r}="mensual",0.25,""))))'
    t[f"I{r}"] = f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r}/60*Supuestos!$B$4)'
    # Costo por hora = sueldo × carga ÷ (horas/semana × 4.33)
    t[f"J{r}"] = f'=IF(OR(I{r}="",H{r}=""),"",I{r}*(H{r}*Supuestos!$B$3/(Supuestos!$B$2*4.33)))'
    t[f"K{r}"] = f'=IF(I{r}="","",IF(AND(G{r}="A",I{r}>=Supuestos!$B$8),"SÍ","no"))'

dv_freq = DataValidation(type="list", formula1='"diaria,semanal,quincenal,mensual"', allow_blank=True)
dv_tipo = DataValidation(type="list", formula1='"A,B"', allow_blank=True)
dv_paso = DataValidation(type="list", formula1='"habilidad,conector,proyecto,prompt"', allow_blank=True)
for dv, col in ((dv_freq, "D"), (dv_tipo, "G"), (dv_paso, "L")):
    t.add_data_validation(dv)
    dv.add(f"{col}2:{col}{N + 1}")

# Ejemplo
t["A2"], t["B2"], t["C2"], t["D2"], t["F2"], t["G2"], t["H2"], t["L2"], t["M2"] = (
    "Finanzas", "Conciliar el banco", "Laura", "semanal", 180, "A", 18000, "habilidad", "Contralor")

# --- Resumen ---------------------------------------------------------------------
res = wb.create_sheet("Resumen por área")
res.append(["Área", "Tareas", "Horas al año tipo A", "Costo al año tipo A", "Candidatas", "Horas recuperables (×realización)"])
areas = ["Dirección", "Operaciones", "Finanzas", "RH", "Marketing", "Ventas", "TI", "Administración", "Atención a clientes"]
for i, a in enumerate(areas, start=2):
    res[f"A{i}"] = a
    res[f"B{i}"] = f'=COUNTIF(Tareas!$A$2:$A${N+1},A{i})'
    res[f"C{i}"] = f'=SUMIFS(Tareas!$I$2:$I${N+1},Tareas!$A$2:$A${N+1},A{i},Tareas!$G$2:$G${N+1},"A")'
    res[f"D{i}"] = f'=SUMIFS(Tareas!$J$2:$J${N+1},Tareas!$A$2:$A${N+1},A{i},Tareas!$G$2:$G${N+1},"A")'
    res[f"E{i}"] = f'=COUNTIFS(Tareas!$A$2:$A${N+1},A{i},Tareas!$K$2:$K${N+1},"SÍ")'
    res[f"F{i}"] = f'=C{i}*Supuestos!$B$5'
fin = len(areas) + 2
res[f"A{fin}"] = "TOTAL"
for col in "BCDEF":
    res[f"{col}{fin}"] = f"=SUM({col}2:{col}{fin-1})"
for i, w in enumerate([22, 10, 20, 20, 12, 30], start=1):
    res.column_dimensions[get_column_letter(i)].width = w

# --- Estilo de encabezados ---------------------------------------------------------
for hoja in (sup, t, res):
    for celda in hoja[1]:
        celda.font = Font(bold=True, color=TEXTO)
        celda.fill = PatternFill("solid", fgColor=NAVY)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.freeze_panes = "A2"
for r in range(2, N + 2):
    t[f"J{r}"].number_format = '"$"#,##0'
    t[f"I{r}"].number_format = "0.0"
for i in range(2, fin + 1):
    res[f"D{i}"].number_format = '"$"#,##0'
    res[f"C{i}"].number_format = res[f"F{i}"].number_format = "0.0"

wb.move_sheet("Tareas", offset=-1)
destino = AQUI / "inventario-tareas.xlsx"
wb.save(destino)
print(f"  ✓ {destino.name}")
