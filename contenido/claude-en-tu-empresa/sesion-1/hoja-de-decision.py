#!/usr/bin/env python
"""
hoja-de-decision.py — Genera la hoja de decisión de la sesión 1 (Total Coach) como .xlsx
con fórmulas vivas: un proyecto por columna, un criterio por fila, suma ponderada al pie.

Los pesos son los del Excel real de Total Coach (20/15/10/30/25), que son los que producen
los totales del ejemplo (8.0 · 8.4 · 8.4 · 7.8). La lámina 26 del deck original decía otro
orden (30/25/20/15/10): ese era un error de rotulado, no del método.

    python hoja-de-decision.py [destino.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")

NAVY, CYAN, DEEP, SOFT, LIME, SURFACE, LINE, TEXT = "0A1A2F", "00A0DB", "006E96", "DFF4FC", "C6F24E", "EAF4FA", "D2E3EC", "072835"
destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("hoja-de-decision.xlsx")

wb = Workbook()

# ── Hoja 1 · Decisión ───────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Decisión"
fino = Side(style="thin", color=LINE)
borde = Border(left=fino, right=fino, top=fino, bottom=fino)
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
izq = Alignment(horizontal="left", vertical="center", wrap_text=True)

ws["B1"] = "Hoja de decisión · Claude en tu Empresa"
ws["B1"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
ws["B2"] = "Un proyecto por columna. Califica del 1 al 10 cada criterio. La hoja calcula la suma ponderada."
ws["B2"].font = Font(name="Arial", size=10, color="4E6572")

CRITERIOS = [
    ("Conocimiento técnico y reto de implementación", 0.20, "¿Tenemos la capacidad interna? 10 = sí, sin ayuda"),
    ("Facilidad y tiempo / complejidad", 0.15, "¿Cuánto esfuerzo requiere? 10 = poco y rápido"),
    ("Probabilidad de éxito", 0.10, "¿Qué tan viable es lograrlo? 10 = casi seguro"),
    ("Impacto en la empresa: tiempo, costo, control", 0.30, "¿Cuánto mueve el negocio? 10 = mucho"),
    ("ROI / presupuesto", 0.25, "¿Qué retorno genera contra la inversión? 10 = alto"),
]
PROYECTOS = ["Cobranza con IA", "Remarketing de clientes con IA", "Prospección de clientes nuevos con IA", "Análisis estadístico de pólizas canceladas"]
EJEMPLO = [[10, 4, 7, 8, 9], [8, 4, 7, 10, 10], [6, 6, 8, 10, 10], [9, 8, 6, 9, 6]]
INVERSION = [1500, 3500, 3500, 7000]
HORAS = [200, 100, 100, 500]

fila_cab = 4
ws.cell(fila_cab, 2, "Peso").font = Font(name="Arial", bold=True, color="FFFFFF")
ws.cell(fila_cab, 3, "Criterio").font = Font(name="Arial", bold=True, color="FFFFFF")
ws.cell(fila_cab, 4, "Cómo calificar").font = Font(name="Arial", bold=True, color="FFFFFF")
for c in (2, 3, 4):
    ws.cell(fila_cab, c).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(fila_cab, c).alignment = centro
    ws.cell(fila_cab, c).border = borde
for k, nombre in enumerate(PROYECTOS):
    celda = ws.cell(fila_cab, 5 + k, nombre)
    celda.font = Font(name="Arial", bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=DEEP)
    celda.alignment = centro
    celda.border = borde
ws.row_dimensions[fila_cab].height = 48

primera = fila_cab + 1
for i, (crit, peso, ayuda) in enumerate(CRITERIOS):
    f = primera + i
    ws.cell(f, 2, peso).number_format = "0%"
    ws.cell(f, 3, crit)
    ws.cell(f, 4, ayuda).font = Font(name="Arial", size=9, color="4E6572")
    for c in (2, 3, 4):
        ws.cell(f, c).border = borde
        ws.cell(f, c).alignment = izq if c != 2 else centro
    ws.cell(f, 2).fill = PatternFill("solid", fgColor=SOFT)
    ws.cell(f, 2).font = Font(name="Arial", bold=True, color=DEEP)
    for k in range(len(PROYECTOS)):
        celda = ws.cell(f, 5 + k, EJEMPLO[k][i])
        celda.alignment = centro
        celda.border = borde
        celda.font = Font(name="Arial", size=12)
ultima = primera + len(CRITERIOS) - 1

f_inv, f_hrs = ultima + 1, ultima + 2
ws.cell(f_inv, 3, "Monto de inversión estimado (MXN)")
ws.cell(f_hrs, 3, "Horas de programación o configuración")
for f, valores, fmt in ((f_inv, INVERSION, '"$"#,##0'), (f_hrs, HORAS, "#,##0")):
    for c in (2, 3, 4):
        ws.cell(f, c).border = borde
        ws.cell(f, c).fill = PatternFill("solid", fgColor=SURFACE)
    ws.cell(f, 3).font = Font(name="Arial", italic=True, color=TEXT)
    for k, v in enumerate(valores):
        celda = ws.cell(f, 5 + k, v)
        celda.number_format = fmt
        celda.alignment = centro
        celda.border = borde
        celda.fill = PatternFill("solid", fgColor=SURFACE)

f_tot = f_hrs + 2
ws.cell(f_tot, 2, f"=SUM(B{primera}:B{ultima})").number_format = "0%"
ws.cell(f_tot, 3, "Calificación ponderada").font = Font(name="Arial", bold=True, size=12, color=NAVY)
for c in (2, 3, 4):
    ws.cell(f_tot, c).border = borde
    ws.cell(f_tot, c).fill = PatternFill("solid", fgColor=LIME)
ws.cell(f_tot, 2).font = Font(name="Arial", bold=True, color=NAVY)
ws.cell(f_tot, 2).alignment = centro
for k in range(len(PROYECTOS)):
    col = get_column_letter(5 + k)
    celda = ws.cell(f_tot, 5 + k, f"=SUMPRODUCT($B${primera}:$B${ultima},{col}{primera}:{col}{ultima})")
    celda.number_format = "0.0"
    celda.font = Font(name="Arial", bold=True, size=14, color=NAVY)
    celda.fill = PatternFill("solid", fgColor=LIME)
    celda.alignment = centro
    celda.border = borde
ws.cell(f_tot + 1, 3, "Si los pesos no suman 100%, la celda de la izquierda lo avisa.").font = Font(name="Arial", size=9, color="4E6572")
ws.cell(f_tot + 2, 3, "Empate: revisa inversión, horas, capacidad del equipo y urgencia. El número ordena; no decide solo.").font = Font(name="Arial", size=9, color="4E6572")

dv = DataValidation(type="whole", operator="between", formula1=1, formula2=10, showErrorMessage=True,
                    errorTitle="Del 1 al 10", error="Califica cada criterio con un entero del 1 al 10.")
ws.add_data_validation(dv)
dv.add(f"E{primera}:{get_column_letter(4 + len(PROYECTOS))}{ultima}")

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 9
ws.column_dimensions["C"].width = 44
ws.column_dimensions["D"].width = 40
for k in range(len(PROYECTOS)):
    ws.column_dimensions[get_column_letter(5 + k)].width = 22
ws.freeze_panes = "E5"

# ── Hoja 2 · Cómo se usa ────────────────────────────────────────────────────────
ins = wb.create_sheet("Cómo se usa")
pasos = [
    "1. Escribe el nombre de cada proyecto en la fila 4, uno por columna. Borra los cuatro de ejemplo.",
    "2. Califica cada criterio del 1 al 10, en consenso con tu equipo. No en solitario.",
    "3. Anota el monto de inversión estimado y las horas de programación o configuración.",
    "4. Lee la calificación ponderada al pie. Ordena, pero no decide sola.",
    "5. Si dos proyectos empatan, decide con inversión, horas, capacidad del equipo y urgencia.",
    "6. Elijan dos proyectos, máximo tres. Los demás van al backlog del siguiente ciclo.",
    "",
    "Los pesos (20 · 15 · 10 · 30 · 25) son los que usa Total Coach. Impacto y retorno concentran el 55%:",
    "la facilidad no gana sola, debe sostenerse con valor real de negocio.",
]
for i, p in enumerate(pasos, start=2):
    ins.cell(i, 2, p).font = Font(name="Arial", size=11 if p and p[0].isdigit() else 10, color=TEXT if p and p[0].isdigit() else "4E6572")
ins.cell(1, 2, "Cómo se usa la hoja de decisión").font = Font(name="Arial", size=14, bold=True, color=NAVY)
ins.column_dimensions["B"].width = 110

destino.parent.mkdir(parents=True, exist_ok=True)
wb.save(destino)
print(f"  ✓ {destino}  ·  totales de ejemplo esperados: 8.0 · 8.4 · 8.4 · 7.8")
